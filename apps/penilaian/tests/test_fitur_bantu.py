import datetime
import io

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.akun.peran import Peran
from apps.dokumen.models import StatusDokumen
from apps.dokumen.tests.factories import BuktiEPFactory, DokumenFactory
from apps.penilaian import services
from apps.penilaian.capaian import ep_aktif, kekurangan_bukti
from apps.penilaian.models import PeriodeAkreditasi
from apps.penilaian.tugas import langkah_persiapan, tugas_untuk
from apps.pps.models import RencanaPerbaikan
from apps.standar.tests.factories import BabFactory, ElemenPenilaianFactory

pytestmark = pytest.mark.django_db


@pytest.fixture
def periode():
    return PeriodeAkreditasi.objects.create(
        nama="Periode contoh", tanggal_mulai=datetime.date(2026, 1, 1), aktif=True
    )


@pytest.fixture
def bab(pokja):
    return BabFactory(pokja=pokja)


def _judul(tugas):
    return {t.judul: t.jumlah for t in tugas}


def test_kekurangan_membedakan_belum_ada_dan_menunggu_verifikasi(bab):
    ep = ElemenPenilaianFactory(standar__bab=bab, jenis_bukti=["R", "D", "O"])
    BuktiEPFactory(elemen_penilaian=ep, jenis_bukti="R", dokumen__status=StatusDokumen.DISETUJUI)
    BuktiEPFactory(elemen_penilaian=ep, jenis_bukti="D", dokumen__status=StatusDokumen.DIAJUKAN)
    lengkap = ElemenPenilaianFactory(standar__bab=bab, jenis_bukti=["R"])
    BuktiEPFactory(
        elemen_penilaian=lengkap, jenis_bukti="R", dokumen__status=StatusDokumen.DISETUJUI
    )

    (k,) = kekurangan_bukti(ep_aktif([bab.pk]))
    assert k.ep == ep
    assert (k.kurang, k.proses, k.belum_ada) == (["D", "O"], ["D"], ["O"])


def test_tugas_ketua_pokja(buat_user, pokja, bab, periode):
    ketua = buat_user(Peran.KETUA_POKJA, pokja=[pokja])
    ElemenPenilaianFactory(standar__bab=bab, jenis_bukti=["R"])
    DokumenFactory(pokja=pokja)  # draft siap diajukan
    DokumenFactory(pokja=pokja, status=StatusDokumen.DIKEMBALIKAN)
    DokumenFactory(pokja=pokja, status=StatusDokumen.KEDALUWARSA)
    DokumenFactory(status=StatusDokumen.DIKEMBALIKAN)  # pokja lain, tidak dihitung
    tugas = _judul(tugas_untuk(ketua, periode))
    assert tugas == {
        "Dokumen dikembalikan, perlu diperbaiki": 1,
        "Dokumen kedaluwarsa, perlu direview": 1,
        "Draft dokumen siap diajukan": 1,
        "EP pokja Anda yang buktinya belum lengkap": 1,
        "EP belum dinilai pada periode ini": 1,
    }


def test_tugas_anggota_tanpa_tugas_ketua(buat_user, pokja, bab, periode):
    anggota = buat_user(Peran.ANGGOTA_POKJA, pokja=[pokja])
    DokumenFactory(pokja=pokja)
    tugas = _judul(tugas_untuk(anggota, periode))
    assert "Draft dokumen siap diajukan" not in tugas
    assert "EP belum dinilai pada periode ini" not in tugas


def test_tugas_verifikator_dan_pj_pps(buat_user, pokja, periode):
    komite = buat_user(Peran.KOMITE_MUTU)
    direktur = buat_user(Peran.DIREKTUR)
    DokumenFactory(pokja=pokja, status=StatusDokumen.DIAJUKAN)
    DokumenFactory(pokja=pokja, status=StatusDokumen.DIVERIFIKASI_KOMITE)  # SPO → direktur
    assert _judul(tugas_untuk(komite, periode)) == {
        "Dokumen menunggu verifikasi/persetujuan Anda": 1
    }
    assert _judul(tugas_untuk(direktur, periode)) == {
        "Dokumen menunggu verifikasi/persetujuan Anda": 1
    }
    RencanaPerbaikan.objects.create(
        periode=periode,
        elemen_penilaian=ElemenPenilaianFactory(),
        rekomendasi="Contoh",
        pj=direktur,
        dibuat_oleh=komite,
        tenggat=datetime.date(2020, 1, 1),
    )
    t = next(t for t in tugas_untuk(direktur, periode) if "PPS" in t.judul)
    assert t.penting and "lewat tenggat" in t.keterangan


def test_langkah_persiapan_hanya_pengelola(buat_user, periode):
    assert langkah_persiapan(buat_user(Peran.KETUA_POKJA)) == []
    langkah = langkah_persiapan(buat_user(Peran.KOMITE_MUTU))
    assert [lk.selesai for lk in langkah] == [False, True, False, True, False]


def test_dashboard_menampilkan_tugas(client, buat_user, pokja, periode):
    DokumenFactory(pokja=pokja, status=StatusDokumen.DIKEMBALIKAN)
    client.force_login(buat_user(Peran.ANGGOTA_POKJA, pokja=[pokja]))
    resp = client.get(reverse("beranda"))
    assert "Dokumen dikembalikan, perlu diperbaiki" in resp.text
    assert "persiapan" not in resp.context


def test_detail_ep(client, buat_user, pokja, bab, periode):
    ep = ElemenPenilaianFactory(standar__bab=bab, jenis_bukti=["R", "D"])
    BuktiEPFactory(
        elemen_penilaian=ep,
        jenis_bukti="R",
        dokumen__judul="Kebijakan contoh X",
        dokumen__status=StatusDokumen.DISETUJUI,
    )
    ketua = buat_user(Peran.KETUA_POKJA, pokja=[pokja])
    services.nilai_ep(user=ketua, periode=periode, ep=ep, nilai=5)
    client.force_login(ketua)
    resp = client.get(reverse("penilaian:ep", args=[ep.pk]))
    assert resp.status_code == 200
    assert "Kebijakan contoh X" in resp.text
    assert [(k, len(b)) for k, _, b in resp.context["bukti_per_jenis"]] == [("R", 1), ("D", 0)]
    assert resp.context["riwayat"][0].skor == 5 and resp.context["bisa_nilai"]
    lain = ElemenPenilaianFactory()
    assert client.get(reverse("penilaian:ep", args=[lain.pk])).status_code == 404


def test_halaman_kekurangan_dan_excel(client, buat_user, pokja, bab):
    ep = ElemenPenilaianFactory(standar__bab=bab, jenis_bukti=["R"])
    ElemenPenilaianFactory(jenis_bukti=["R"])  # bab lain
    client.force_login(buat_user(Peran.ANGGOTA_POKJA, pokja=[pokja]))
    resp = client.get(reverse("penilaian:kekurangan"))
    assert [k.ep for k in resp.context["halaman"]] == [ep]
    assert client.get(reverse("penilaian:kekurangan"), {"bab": BabFactory().pk}).status_code == 404
    resp = client.get(reverse("penilaian:kekurangan"), {"format": "xlsx"})
    baris = list(load_workbook(io.BytesIO(resp.content)).active.values)
    assert len(baris) == 2 and baris[1][3] == ep.kode


def test_laporan_excel_penilaian(client, buat_user, pokja, bab, periode):
    komite = buat_user(Peran.KOMITE_MUTU)
    ep = ElemenPenilaianFactory(standar__bab=bab)
    services.nilai_ep(user=komite, periode=periode, ep=ep, nilai="tdd", fakta_analisis="Fakta X")
    client.force_login(komite)
    wb = load_workbook(io.BytesIO(client.get(reverse("penilaian:laporan")).content))
    assert wb.sheetnames == ["Rekap Bab", "Detail EP"]
    rekap = list(wb["Rekap Bab"].values)
    assert rekap[-1][0] == "TOTAL" and rekap[1][5] == 1  # 1 TDD
    detail = list(wb["Detail EP"].values)
    assert detail[1][7] == "TDD" and detail[1][8] == "Fakta X"


def test_cari_tersaring_akses(client, buat_user, pokja, bab):
    milik = ElemenPenilaianFactory(standar__bab=bab, uraian="Uraian EP kata-unik")
    ElemenPenilaianFactory(uraian="Uraian EP kata-unik lain")
    dok = DokumenFactory(pokja=pokja, judul="SPO kata-unik")
    DokumenFactory(judul="SPO kata-unik pokja lain")
    client.force_login(buat_user(Peran.ANGGOTA_POKJA, pokja=[pokja]))
    resp = client.get(reverse("cari"), {"q": "kata-unik"})
    assert list(resp.context["daftar_ep"]) == [milik]
    assert list(resp.context["daftar_dokumen"]) == [dok]
