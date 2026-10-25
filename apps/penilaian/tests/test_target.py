import datetime
import io

import pytest
from django.core.exceptions import PermissionDenied
from django.core.management import call_command
from django.urls import reverse
from openpyxl import load_workbook

from apps.akun.peran import Peran
from apps.dokumen.models import StatusDokumen
from apps.dokumen.tests.factories import BuktiEPFactory
from apps.penilaian import services
from apps.penilaian import target as target_svc
from apps.penilaian.models import PeriodeAkreditasi, TargetPokja
from apps.penilaian.tugas import tugas_untuk
from apps.standar.tests.factories import BabFactory, ElemenPenilaianFactory
from apps.verifikasi.models import Notifikasi

pytestmark = pytest.mark.django_db
H = datetime.date(2026, 6, 1)  # "hari ini" dalam test
S = TargetPokja.Status


@pytest.fixture
def periode():
    return PeriodeAkreditasi.objects.create(
        nama="Periode contoh", tanggal_mulai=datetime.date(2026, 1, 1), aktif=True
    )


@pytest.fixture
def komite(buat_user):
    return buat_user(Peran.KOMITE_MUTU)


@pytest.fixture
def ep(pokja):
    """Bab pokja berisi 4 EP; 1 lengkap bukti (25%), 2 dinilai (50%)."""
    bab = BabFactory(pokja=pokja)
    daftar = [ElemenPenilaianFactory(standar__bab=bab, jenis_bukti=["R"]) for _ in range(4)]
    BuktiEPFactory(
        elemen_penilaian=daftar[0], jenis_bukti="R", dokumen__status=StatusDokumen.DISETUJUI
    )
    return daftar


def _target(komite, periode, pokja, ukuran, nilai, mulai, tenggat):
    target_svc.buat_target_banyak(
        user=komite,
        periode=periode,
        daftar_pokja=[pokja],
        ukuran=ukuran,
        nilai_target=nilai,
        tanggal_mulai=mulai,
        tenggat=tenggat,
    )
    return TargetPokja.objects.get(pokja=pokja, ukuran=ukuran, tenggat=tenggat)


def test_harapan_linier():
    t = TargetPokja(
        nilai_target=100,
        tanggal_mulai=datetime.date(2026, 1, 1),
        tenggat=datetime.date(2026, 1, 11),
    )
    assert target_svc.harapan_pada(t, datetime.date(2025, 12, 1)) == 0
    assert target_svc.harapan_pada(t, datetime.date(2026, 1, 6)) == 50
    assert target_svc.harapan_pada(t, datetime.date(2026, 2, 1)) == 100


@pytest.mark.parametrize(
    ("aktual", "harapan", "tenggat", "status"),
    [
        (80, 50, H, S.TERCAPAI),
        (60, 50, H, S.SESUAI),
        (40, 50, H, S.BERISIKO),
        (40, 100, H - datetime.timedelta(days=1), S.TERLAMBAT),
        (80, 100, H - datetime.timedelta(days=1), S.TERCAPAI),
    ],
)
def test_tentukan_status(aktual, harapan, tenggat, status):
    assert (
        target_svc.tentukan_status(
            nilai_target=80, aktual=aktual, harapan=harapan, hari_ini=H, tenggat=tenggat
        )
        == status
    )


def test_evaluasi_memakai_capaian_pokja(komite, periode, pokja, ep):
    services.nilai_ep(user=komite, periode=periode, ep=ep[0], nilai=10)
    services.nilai_ep(user=komite, periode=periode, ep=ep[1], nilai=0)
    U = TargetPokja.Ukuran
    _target(
        komite,
        periode,
        pokja,
        U.KELENGKAPAN_BUKTI,
        100,
        datetime.date(2026, 5, 1),
        datetime.date(2026, 7, 1),
    )
    _target(
        komite,
        periode,
        pokja,
        U.PROGRES_PENILAIAN,
        50,
        datetime.date(2026, 5, 1),
        datetime.date(2026, 7, 1),
    )
    _target(
        komite,
        periode,
        pokja,
        U.CAPAIAN_SKOR,
        80,
        datetime.date(2026, 1, 1),
        datetime.date(2026, 5, 1),
    )
    hasil = {h.target.ukuran: h for h in target_svc.evaluasi(TargetPokja.objects.all(), hari_ini=H)}
    lengkap = hasil[U.KELENGKAPAN_BUKTI]
    assert lengkap.aktual == 25.0 and lengkap.harapan == pytest.approx(50.8, abs=0.1)
    assert lengkap.status == S.BERISIKO
    assert hasil[U.PROGRES_PENILAIAN].status == S.TERCAPAI  # 2/4 = 50%
    skor = hasil[U.CAPAIAN_SKOR]
    assert (skor.aktual, skor.status, skor.hari_lewat) == (50.0, S.TERLAMBAT, 31)


def test_buat_banyak_dan_hak_akses(buat_user, komite, periode, pokja):
    from apps.akun.tests.factories import PokjaFactory

    lain = PokjaFactory()
    args = dict(
        periode=periode,
        daftar_pokja=[pokja, lain],
        ukuran=TargetPokja.Ukuran.KELENGKAPAN_BUKTI,
        nilai_target=100,
        tenggat=datetime.date(2026, 9, 1),
    )
    assert target_svc.buat_target_banyak(user=komite, **args) == (2, 0)
    assert target_svc.buat_target_banyak(user=komite, **{**args, "nilai_target": 90}) == (0, 2)
    assert set(TargetPokja.objects.values_list("nilai_target", flat=True)) == {90}
    with pytest.raises(PermissionDenied):
        target_svc.buat_target_banyak(user=buat_user(Peran.DIREKTUR), **args)


def test_halaman_target_per_peran(client, buat_user, komite, periode, pokja, ep):
    from apps.akun.tests.factories import PokjaFactory

    milik = _target(
        komite,
        periode,
        pokja,
        TargetPokja.Ukuran.KELENGKAPAN_BUKTI,
        100,
        datetime.date(2026, 1, 1),
        datetime.date(2030, 1, 1),
    )
    _target(
        komite,
        periode,
        PokjaFactory(),
        TargetPokja.Ukuran.KELENGKAPAN_BUKTI,
        100,
        datetime.date(2026, 1, 1),
        datetime.date(2030, 1, 1),
    )
    url = reverse("penilaian:target_list")

    client.force_login(buat_user(Peran.DIREKTUR))
    assert len(client.get(url).context["daftar"]) == 2
    client.force_login(buat_user(Peran.ANGGOTA_POKJA, pokja=[pokja]))
    resp = client.get(url)
    assert [h.target for h in resp.context["daftar"]] == [milik]
    assert not resp.context["bisa_kelola"]
    assert client.get(reverse("penilaian:target_baru")).status_code == 403
    client.force_login(buat_user(Peran.TAMU_SURVEIOR))
    assert client.get(url).context["daftar"] == []


def test_form_tetapkan_target_banyak_pokja(client, komite, periode, pokja):
    from apps.akun.tests.factories import PokjaFactory

    lain = PokjaFactory()
    client.force_login(komite)
    data = {
        "periode": periode.pk,
        "daftar_pokja": [pokja.pk, lain.pk],
        "ukuran": "kelengkapan_bukti",
        "nilai_target": "100",
        "tenggat": "2026-11-30",
    }
    assert client.post(reverse("penilaian:target_baru"), data).status_code == 302
    assert TargetPokja.objects.count() == 2
    resp = client.post(reverse("penilaian:target_baru"), {**data, "tenggat": "2025-12-01"})
    assert "tenggat" in resp.context["form"].errors
    t = TargetPokja.objects.first()
    client.post(reverse("penilaian:target_hapus", args=[t.pk]))
    assert TargetPokja.objects.count() == 1


def test_dashboard_tugas_dan_excel(client, buat_user, komite, periode, pokja, ep):
    _target(
        komite,
        periode,
        pokja,
        TargetPokja.Ukuran.KELENGKAPAN_BUKTI,
        100,
        datetime.date(2026, 1, 1),
        datetime.date(2026, 2, 1),
    )  # sudah lewat → terlambat
    ketua = buat_user(Peran.KETUA_POKJA, pokja=[pokja])
    judul = {t.judul for t in tugas_untuk(ketua, periode)}
    assert "Target pokja Anda tertinggal dari jadwal" in judul
    assert "Target pokja tertinggal" in {
        t.judul for t in tugas_untuk(buat_user(Peran.DIREKTUR), periode)
    }

    client.force_login(buat_user(Peran.DIREKTUR))
    resp = client.get(reverse("beranda"))
    assert resp.context["target_tertinggal"][0].status == S.TERLAMBAT
    wb = load_workbook(io.BytesIO(client.get(reverse("penilaian:laporan")).content))
    assert "Target Pokja" in wb.sheetnames
    assert list(wb["Target Pokja"].values)[1][8] == "Terlambat"


def test_notifikasi_hanya_saat_status_berubah(buat_user, komite, periode, pokja, ep):
    ketua = buat_user(Peran.KETUA_POKJA, pokja=[pokja])
    t = _target(
        komite,
        periode,
        pokja,
        TargetPokja.Ukuran.KELENGKAPAN_BUKTI,
        100,
        datetime.date(2026, 1, 1),
        datetime.date(2026, 2, 1),
    )
    call_command("kirim_pengingat", verbosity=0)
    call_command("kirim_pengingat", verbosity=0)
    assert Notifikasi.objects.filter(penerima=ketua, judul__contains="terlambat").count() == 1
    assert Notifikasi.objects.filter(penerima=komite).count() == 1  # terlambat → komite ikut
    t.refresh_from_db()
    assert t.status_dinotifikasi == S.TERLAMBAT
