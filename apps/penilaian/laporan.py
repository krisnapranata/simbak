"""Laporan Excel: rekap penilaian mandiri dan daftar kekurangan bukti."""

from django.db.models import Prefetch
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from apps.standar.models import JenisBukti

from .capaian import bukti_disetujui, capaian_per_bab, ep_aktif, kekurangan_bukti, total
from .models import PenilaianEP

TEBAL = Font(bold=True)
ISI_HEADER = PatternFill("solid", fgColor="E2E8F0")


def _sheet(wb, judul, kolom, pertama=False):
    ws = wb.active if pertama else wb.create_sheet()
    ws.title = judul
    ws.append([k for k, _ in kolom])
    for sel, (_, lebar) in zip(ws[1], kolom, strict=True):
        sel.font, sel.fill = TEBAL, ISI_HEADER
        ws.column_dimensions[sel.column_letter].width = lebar
    ws.freeze_panes = "A2"
    return ws


def _rapikan(ws):
    for row in ws.iter_rows(min_row=2):
        for sel in row:
            sel.alignment = Alignment(wrap_text=True, vertical="top")


def _persen(nilai):
    return None if nilai is None else round(nilai / 100, 4)


def _label_jenis(kode):
    return ", ".join(f"{k} ({JenisBukti(k).label})" for k in kode)


def laporan_penilaian(*, periode, bab_qs, hasil_target=None) -> Workbook:
    wb = Workbook()
    baris = capaian_per_bab(periode=periode, bab_qs=bab_qs)

    ws = _sheet(
        wb,
        "Rekap Bab",
        [
            ("Bab", 10),
            ("Nama bab", 40),
            ("Pokja", 16),
            ("Jumlah EP", 11),
            ("Dinilai", 9),
            ("TDD", 7),
            ("Total skor", 11),
            ("Capaian", 10),
            ("Bukti lengkap", 13),
        ],
        pertama=True,
    )
    for bab, c in [*baris, (None, total(baris))]:
        ws.append(
            [
                bab.kode if bab else "TOTAL",
                bab.nama if bab else "",
                (bab.pokja.kode if bab and bab.pokja else ""),
                c.jumlah_ep,
                c.dinilai,
                c.tdd,
                c.total_skor,
                _persen(c.persen),
                _persen(c.persen_lengkap),
            ]
        )
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=9):
        for sel in row:
            sel.number_format = "0.0%"
    for sel in ws[ws.max_row]:
        sel.font = TEBAL

    ws = _sheet(
        wb,
        "Detail EP",
        [
            ("Bab", 9),
            ("Standar", 14),
            ("EP", 6),
            ("Uraian EP", 55),
            ("Jenis wajib", 14),
            ("Bukti disetujui", 14),
            ("Kurang", 12),
            ("Skor", 7),
            ("Fakta & analisis", 45),
            ("Catatan", 30),
            ("Dinilai oleh", 20),
        ],
    )
    ep_qs = ep_aktif([b.pk for b, _ in baris])
    tercakup = bukti_disetujui(ep_qs)
    penilaian = PenilaianEP.objects.filter(periode=periode).select_related("dinilai_oleh")
    daftar = (
        ep_qs.select_related("standar__bab")
        .prefetch_related(Prefetch("penilaian", queryset=penilaian, to_attr="nilai_periode"))
        .order_by(
            "standar__bab__kelompok__urutan",
            "standar__bab__urutan",
            "standar__urutan",
            "standar__kode",
            "urutan",
            "kode",
        )
    )
    for ep in daftar:
        ok = tercakup.get(ep.pk, set())
        n = ep.nilai_periode[0] if ep.nilai_periode else None
        ws.append(
            [
                ep.standar.bab.kode,
                ep.standar.kode,
                ep.kode,
                ep.uraian,
                ", ".join(ep.jenis_bukti),
                ", ".join(j for j in ep.jenis_bukti if j in ok),
                ", ".join(j for j in ep.jenis_bukti if j not in ok),
                ("TDD" if n.tdd else n.skor) if n else "belum dinilai",
                n.fakta_analisis if n else "",
                n.catatan if n else "",
                str(n.dinilai_oleh) if n else "",
            ]
        )
    _rapikan(ws)
    if hasil_target:
        _sheet_target(wb, hasil_target)
    return wb


KOLOM_TARGET = [
    ("Pokja", 16),
    ("Ukuran", 34),
    ("Target", 9),
    ("Aktual", 9),
    ("Seharusnya hari ini", 12),
    ("Mulai", 12),
    ("Tenggat", 12),
    ("Sisa hari", 10),
    ("Status", 14),
    ("Keterangan", 40),
]


def _sheet_target(wb, hasil, pertama=False):
    ws = _sheet(wb, "Target Pokja", KOLOM_TARGET, pertama=pertama)
    for h in hasil:
        t = h.target
        ws.append(
            [
                t.pokja.kode,
                t.get_ukuran_display(),
                h.nilai_target / 100,
                h.aktual / 100,
                h.harapan / 100,
                t.mulai,
                t.tenggat,
                h.sisa_hari,
                h.label_status,
                t.keterangan,
            ]
        )
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=5):
        for sel in row:
            sel.number_format = "0.0%"
    _rapikan(ws)
    return ws


def laporan_target(hasil) -> Workbook:
    wb = Workbook()
    _sheet_target(wb, hasil, pertama=True)
    return wb


def laporan_kekurangan(daftar) -> Workbook:
    wb = Workbook()
    ws = _sheet(
        wb,
        "Kekurangan Bukti",
        [
            ("Bab", 9),
            ("Pokja", 16),
            ("Standar", 14),
            ("EP", 6),
            ("Uraian EP", 60),
            ("Belum ada bukti", 24),
            ("Menunggu verifikasi", 24),
        ],
        pertama=True,
    )
    for k in daftar:
        bab = k.ep.standar.bab
        ws.append(
            [
                bab.kode,
                bab.pokja.kode if bab.pokja else "",
                k.ep.standar.kode,
                k.ep.kode,
                k.ep.uraian,
                _label_jenis(k.belum_ada),
                _label_jenis(k.proses),
            ]
        )
    _rapikan(ws)
    return wb


def kekurangan_untuk(bab_qs):
    return kekurangan_bukti(ep_aktif(list(bab_qs.values_list("id", flat=True))))
