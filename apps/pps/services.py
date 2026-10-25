"""Logika bisnis Perencanaan Perbaikan Strategis (PPS)."""

from django.core.exceptions import PermissionDenied
from django.db import transaction
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from apps.inti.services import hapus_lunak
from apps.penilaian.models import PenilaianEP, SkorEP

from .models import RencanaPerbaikan
from .permissions import (
    bisa_buat_pps_otomatis,
    bisa_edit_pps,
    bisa_kelola_pps_bab,
    bisa_update_progres,
)

FIELD_PROGRES = ["status", "catatan_progres"]


def simpan_pps(*, user, pps: RencanaPerbaikan, field_diubah=None) -> RencanaPerbaikan:
    """Buat/ubah PPS. Penanggung jawab non-pengelola hanya boleh mengubah status & progres."""
    if pps.pk is None:
        if not bisa_kelola_pps_bab(user, pps.elemen_penilaian.standar.bab):
            raise PermissionDenied
        pps.dibuat_oleh = user
    elif not bisa_edit_pps(user, RencanaPerbaikan.objects.get(pk=pps.pk)):
        if not bisa_update_progres(user, pps) or set(field_diubah or []) - set(FIELD_PROGRES):
            raise PermissionDenied
    pps.save()
    return pps


def hapus_pps(*, user, pps: RencanaPerbaikan):
    if not bisa_edit_pps(user, pps):
        raise PermissionDenied
    return hapus_lunak(obj=pps)


@transaction.atomic
def buat_dari_penilaian(*, user, periode) -> int:
    """Buat draft PPS untuk EP berskor 0/5 pada periode yang belum punya PPS."""
    if not bisa_buat_pps_otomatis(user):
        raise PermissionDenied
    sudah = RencanaPerbaikan.objects.filter(periode=periode).values("elemen_penilaian_id")
    penilaian = (
        PenilaianEP.objects.filter(
            periode=periode,
            tdd=False,
            skor__in=[SkorEP.TIDAK_TERCAPAI, SkorEP.TERCAPAI_SEBAGIAN],
            elemen_penilaian__dihapus_pada__isnull=True,
        )
        .exclude(elemen_penilaian_id__in=sudah)
        .select_related("elemen_penilaian")
    )
    baru = [
        RencanaPerbaikan(
            periode=periode,
            elemen_penilaian=p.elemen_penilaian,
            rekomendasi=f"(Skor penilaian {p.skor}. Lengkapi dengan rekomendasi surveior.)",
            dibuat_oleh=user,
        )
        for p in penilaian
    ]
    RencanaPerbaikan.objects.bulk_create(baru)
    return len(baru)


KOLOM_EKSPOR = [
    ("Bab", 10),
    ("Standar", 14),
    ("EP", 8),
    ("Rekomendasi", 45),
    ("Rencana kegiatan", 45),
    ("Indikator", 30),
    ("Penanggung jawab", 22),
    ("Tenggat", 12),
    ("Status", 12),
    ("Progres", 40),
]


def ekspor_xlsx(qs) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "PPS"
    ws.append([k for k, _ in KOLOM_EKSPOR])
    for sel, (_, lebar) in zip(ws[1], KOLOM_EKSPOR, strict=True):
        sel.font = Font(bold=True)
        ws.column_dimensions[sel.column_letter].width = lebar
    for p in qs.select_related("elemen_penilaian__standar__bab", "pj"):
        ep = p.elemen_penilaian
        ws.append(
            [
                ep.standar.bab.kode,
                ep.standar.kode,
                ep.kode,
                p.rekomendasi,
                p.rencana_kegiatan,
                p.indikator,
                str(p.pj or ""),
                p.tenggat,
                p.get_status_display(),
                p.catatan_progres,
            ]
        )
    for row in ws.iter_rows(min_row=2):
        for sel in row:
            sel.alignment = Alignment(wrap_text=True, vertical="top")
    return wb
