"""Logika bisnis master standar: impor dari Excel dan pembuatan template."""

import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from django.core.exceptions import PermissionDenied
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.exceptions import InvalidFileException

from apps.akun.models import Pokja

from .models import Bab, ElemenPenilaian, JenisBukti, KelompokStandar, LogImporStandar, Standar
from .permissions import bisa_impor_standar, bisa_proses_log_impor

SHEET_BAB = "Bab"
SHEET_EP = "EP"
KOLOM_BAB = ["kode_kelompok", "nama_kelompok", "kode_bab", "nama_bab", "kode_pokja", "urutan"]
WAJIB_BAB = {"kode_kelompok", "nama_kelompok", "kode_bab", "nama_bab"}
KOLOM_EP = [
    "kode_bab",
    "kode_standar",
    "uraian_standar",
    "maksud_tujuan",
    "kode_ep",
    "uraian_ep",
    "jenis_bukti",
    "urutan",
]
WAJIB_EP = {"kode_bab", "kode_standar", "kode_ep", "uraian_ep", "jenis_bukti"}
ENTITAS = ("kelompok", "pokja", "bab", "standar", "ep")


class GalatImpor(Exception):
    pass


@dataclass
class HasilImpor:
    ringkasan: dict[str, dict[str, int]] = field(
        default_factory=lambda: {e: {"dibuat": 0, "diperbarui": 0, "tetap": 0} for e in ENTITAS}
    )
    dinonaktifkan: list[str] = field(default_factory=list)
    galat: list[str] = field(default_factory=list)
    dry_run: bool = False

    @property
    def berhasil(self) -> bool:
        return not self.galat

    def catat(self, entitas: str, status: str):
        self.ringkasan[entitas][status] += 1


@dataclass
class DataImpor:
    kelompok: dict[str, dict] = field(default_factory=dict)
    bab: dict[str, dict] = field(default_factory=dict)
    standar: dict[tuple[str, str], dict] = field(default_factory=dict)
    ep: dict[tuple[str, str, str], dict] = field(default_factory=dict)


# ---------------------------------------------------------------------------
# Membaca & memvalidasi file
# ---------------------------------------------------------------------------


def _kunci_header(teks) -> str:
    """Normalisasi nama header: abaikan huruf besar/kecil, spasi, dan garis bawah."""
    return re.sub(r"[\s_]+", "", str(teks or "").lower())


def _teks(nilai) -> str:
    if nilai is None:
        return ""
    if isinstance(nilai, float) and nilai.is_integer():
        nilai = int(nilai)
    return str(nilai).strip()


def _cari_sheet(wb, nama):
    for ws in wb.worksheets:
        if ws.title.strip().lower() == nama.lower():
            return ws
    return None


def _baca_sheet(ws, kolom, wajib, galat) -> list[tuple[int, dict[str, str]]]:
    baris = ws.iter_rows(values_only=True)
    header = next(baris, None)
    if header is None:
        galat.append(f"Sheet {ws.title}: kosong.")
        return []
    peta = {}
    target = {_kunci_header(k): k for k in kolom}
    for idx, judul in enumerate(header):
        nama = target.get(_kunci_header(judul))
        if nama and nama not in peta:
            peta[nama] = idx
    hilang = wajib - peta.keys()
    if hilang:
        galat.append(f"Sheet {ws.title}: kolom wajib tidak ada: {', '.join(sorted(hilang))}.")
        return []
    hasil = []
    for nomor, row in enumerate(baris, start=2):
        data = {k: (_teks(row[i]) if i < len(row) else "") for k, i in peta.items()}
        if any(data.values()):
            hasil.append((nomor, data))
    return hasil


def _angka(teks, *, lokasi, galat, default):
    if not teks:
        return default
    try:
        nilai = int(teks)
        if nilai < 0:
            raise ValueError
        return nilai
    except ValueError:
        galat.append(f"{lokasi}: urutan '{teks}' harus bilangan bulat ≥ 0.")
        return default


def _kosong(data, wajib, *, lokasi, galat) -> bool:
    kurang = [k for k in sorted(wajib) if not data.get(k)]
    if kurang:
        galat.append(f"{lokasi}: kolom {', '.join(kurang)} wajib diisi.")
    return bool(kurang)


def _validasi_bab(baris, data: DataImpor, galat):
    for nomor, row in baris:
        lokasi = f"Sheet {SHEET_BAB}, baris {nomor}"
        if _kosong(row, WAJIB_BAB, lokasi=lokasi, galat=galat):
            continue
        kode_kel, kode_bab = row["kode_kelompok"], row["kode_bab"]
        kel = data.kelompok.setdefault(
            kode_kel, {"nama": row["nama_kelompok"], "urutan": len(data.kelompok) + 1}
        )
        if kel["nama"] != row["nama_kelompok"]:
            galat.append(
                f"{lokasi}: nama_kelompok untuk '{kode_kel}' berbeda dengan baris sebelumnya."
            )
        if kode_bab in data.bab:
            galat.append(f"{lokasi}: kode_bab '{kode_bab}' duplikat.")
            continue
        data.bab[kode_bab] = {
            "kode_kelompok": kode_kel,
            "nama": row["nama_bab"],
            "kode_pokja": row.get("kode_pokja", ""),
            "urutan": _angka(
                row.get("urutan"), lokasi=lokasi, galat=galat, default=len(data.bab) + 1
            ),
        }


def _validasi_ep(baris, data: DataImpor, galat):
    bab_db = set(Bab.semua.values_list("kode", flat=True))
    baris_ep: dict[tuple, int] = {}
    baris_standar: dict[tuple, int] = {}
    jumlah_standar_per_bab: dict[str, int] = {}
    jumlah_ep_per_standar: dict[tuple, int] = {}

    for nomor, row in baris:
        lokasi = f"Sheet {SHEET_EP}, baris {nomor}"
        if _kosong(row, WAJIB_EP, lokasi=lokasi, galat=galat):
            continue
        kode_bab, kode_std, kode_ep = row["kode_bab"], row["kode_standar"], row["kode_ep"]

        if kode_bab not in data.bab and kode_bab not in bab_db:
            galat.append(
                f"{lokasi}: kode_bab '{kode_bab}' tidak dikenal (tambahkan di sheet {SHEET_BAB})."
            )
            continue

        jenis = []
        for kode in (k.strip().upper() for k in row["jenis_bukti"].split(",")):
            if not kode:
                continue
            if kode not in JenisBukti.values:
                galat.append(f"{lokasi}: jenis_bukti '{kode}' tidak dikenal.")
            elif kode not in jenis:
                jenis.append(kode)
        if not jenis:
            continue

        kunci_std = (kode_bab, kode_std)
        if kunci_std not in data.standar:
            jumlah_standar_per_bab[kode_bab] = jumlah_standar_per_bab.get(kode_bab, 0) + 1
            data.standar[kunci_std] = {
                "uraian": "",
                "maksud_tujuan": "",
                "urutan": jumlah_standar_per_bab[kode_bab],
            }
            baris_standar[kunci_std] = nomor
        std = data.standar[kunci_std]
        for kolom_file, kolom_model in (
            ("uraian_standar", "uraian"),
            ("maksud_tujuan", "maksud_tujuan"),
        ):
            nilai = row.get(kolom_file, "")
            if not nilai:
                continue
            if not std[kolom_model]:
                std[kolom_model] = nilai
            elif std[kolom_model] != nilai:
                galat.append(
                    f"{lokasi}: {kolom_file} untuk '{kode_std}' berbeda dengan baris sebelumnya."
                )

        kunci_ep = (kode_bab, kode_std, kode_ep)
        if kunci_ep in baris_ep:
            galat.append(
                f"{lokasi}: EP '{kode_ep}' pada standar '{kode_std}' duplikat "
                f"(sudah ada di baris {baris_ep[kunci_ep]})."
            )
            continue
        baris_ep[kunci_ep] = nomor
        jumlah_ep_per_standar[kunci_std] = jumlah_ep_per_standar.get(kunci_std, 0) + 1
        data.ep[kunci_ep] = {
            "uraian": row["uraian_ep"],
            "jenis_bukti": jenis,
            "urutan": _angka(
                row.get("urutan"),
                lokasi=lokasi,
                galat=galat,
                default=jumlah_ep_per_standar[kunci_std],
            ),
        }

    for kunci_std, std in data.standar.items():
        if not std["uraian"]:
            galat.append(
                f"Sheet {SHEET_EP}, baris {baris_standar[kunci_std]}: uraian_standar untuk "
                f"'{kunci_std[1]}' wajib diisi minimal pada satu baris."
            )


def baca_file_impor(path) -> tuple[DataImpor, list[str]]:
    """Baca dan validasi seluruh isi file. Tidak menyentuh database selain membaca."""
    data, galat = DataImpor(), []
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, OSError, KeyError, ValueError):
        return data, ["File tidak dapat dibuka sebagai Excel (.xlsx)."]
    try:
        ws_ep = _cari_sheet(wb, SHEET_EP)
        if ws_ep is None:
            return data, [f"Sheet '{SHEET_EP}' tidak ditemukan."]
        ws_bab = _cari_sheet(wb, SHEET_BAB)
        if ws_bab is not None:
            _validasi_bab(_baca_sheet(ws_bab, KOLOM_BAB, WAJIB_BAB, galat), data, galat)
        _validasi_ep(_baca_sheet(ws_ep, KOLOM_EP, WAJIB_EP, galat), data, galat)
    finally:
        wb.close()
    if not galat and not data.ep:
        galat.append(f"Sheet '{SHEET_EP}' tidak berisi data.")
    return data, galat


# ---------------------------------------------------------------------------
# Menyimpan
# ---------------------------------------------------------------------------


def _simpan(model, *, kunci: dict, nilai: dict, hasil: HasilImpor, entitas: str):
    """update_or_create berbasis kunci alami, termasuk memulihkan baris yang dihapus lunak."""
    obj = model.semua.filter(**kunci).first()
    if obj is None:
        hasil.catat(entitas, "dibuat")
        return model.semua.create(**kunci, **nilai)
    berubah = [f for f, v in nilai.items() if getattr(obj, f) != v]
    for f in berubah:
        setattr(obj, f, nilai[f])
    if obj.dihapus_pada is not None:
        obj.dihapus_pada = None
        berubah.append("dihapus_pada")
    if berubah:
        obj.save(update_fields=[*berubah, "diubah_pada"])
        hasil.catat(entitas, "diperbarui")
    else:
        hasil.catat(entitas, "tetap")
    return obj


def _simpan_data(data: DataImpor, hasil: HasilImpor, *, nonaktifkan_yang_hilang: bool):
    kelompok = {
        kode: _simpan(
            KelompokStandar,
            kunci={"kode": kode},
            nilai={"nama": v["nama"], "urutan": v["urutan"]},
            hasil=hasil,
            entitas="kelompok",
        )
        for kode, v in data.kelompok.items()
    }

    bab = {}
    for kode, v in data.bab.items():
        nilai = {
            "kelompok_id": kelompok[v["kode_kelompok"]].id,
            "nama": v["nama"],
            "urutan": v["urutan"],
        }
        if v["kode_pokja"]:
            pokja = Pokja.semua.filter(kode=v["kode_pokja"]).first()
            if pokja is None:
                pokja = Pokja.semua.create(kode=v["kode_pokja"], nama=f"Pokja {v['nama']}")
                hasil.catat("pokja", "dibuat")
            nilai["pokja_id"] = pokja.id
        bab[kode] = _simpan(Bab, kunci={"kode": kode}, nilai=nilai, hasil=hasil, entitas="bab")
    kode_bab_tambahan = {k for k, _ in data.standar} - bab.keys()
    bab.update({b.kode: b for b in Bab.semua.filter(kode__in=kode_bab_tambahan)})

    standar = {
        (kode_bab, kode): _simpan(
            Standar,
            kunci={"bab_id": bab[kode_bab].id, "kode": kode},
            nilai=v,
            hasil=hasil,
            entitas="standar",
        )
        for (kode_bab, kode), v in data.standar.items()
    }

    id_ep = set()
    for (kode_bab, kode_std, kode), v in data.ep.items():
        ep = _simpan(
            ElemenPenilaian,
            kunci={"standar_id": standar[(kode_bab, kode_std)].id, "kode": kode},
            nilai=v,
            hasil=hasil,
            entitas="ep",
        )
        id_ep.add(ep.id)

    if nonaktifkan_yang_hilang:
        # Hanya bab yang ada di file, agar impor sebagian tidak menonaktifkan bab lain.
        id_bab = [b.id for b in bab.values()]
        sekarang = timezone.now()
        ep_hilang = (
            ElemenPenilaian.objects.filter(standar__bab_id__in=id_bab)
            .exclude(id__in=id_ep)
            .select_related("standar")
        )
        hasil.dinonaktifkan += [f"{ep.standar.kode} EP {ep.kode}" for ep in ep_hilang]
        ep_hilang.update(dihapus_pada=sekarang)
        std_hilang = Standar.objects.filter(bab_id__in=id_bab).exclude(
            id__in=[s.id for s in standar.values()]
        )
        hasil.dinonaktifkan += [f"Standar {s.kode}" for s in std_hilang]
        std_hilang.update(dihapus_pada=sekarang)


def impor_standar(*, path, dry_run=False, nonaktifkan_yang_hilang=False) -> HasilImpor:
    """Impor master standar dari Excel. Bila ada satu saja error, tidak ada yang disimpan."""
    hasil = HasilImpor(dry_run=dry_run)
    data, hasil.galat = baca_file_impor(path)
    if hasil.galat:
        return hasil
    with transaction.atomic():
        _simpan_data(data, hasil, nonaktifkan_yang_hilang=nonaktifkan_yang_hilang)
        if dry_run:
            transaction.set_rollback(True)
    return hasil


# ---------------------------------------------------------------------------
# Alur impor (web & CLI) dengan log
# ---------------------------------------------------------------------------


def _isi_log(log: LogImporStandar, hasil: HasilImpor):
    log.ringkasan = hasil.ringkasan
    log.galat = hasil.galat
    log.dinonaktifkan = hasil.dinonaktifkan


def _hapus_file_sementara(log: LogImporStandar):
    if log.file:
        log.file.delete(save=False)


def buat_pratinjau_impor(*, file, user, nonaktifkan_yang_hilang=False) -> LogImporStandar:
    """Simpan file sementara di storage privat lalu jalankan impor dry-run sebagai pratinjau."""
    if not bisa_impor_standar(user):
        raise PermissionDenied
    log = LogImporStandar(
        diimpor_oleh=user,
        sumber=LogImporStandar.Sumber.WEB,
        nama_file=Path(file.name).name[:255],
        nonaktifkan_yang_hilang=nonaktifkan_yang_hilang,
    )
    log.file.save("impor.xlsx", file, save=False)
    hasil = impor_standar(
        path=log.file.path, dry_run=True, nonaktifkan_yang_hilang=nonaktifkan_yang_hilang
    )
    _isi_log(log, hasil)
    if not hasil.berhasil:
        log.status = LogImporStandar.Status.GAGAL_VALIDASI
        log.selesai_pada = timezone.now()
        _hapus_file_sementara(log)
    log.save()
    return log


def konfirmasi_impor(*, log: LogImporStandar, user) -> LogImporStandar:
    if not bisa_proses_log_impor(user, log):
        raise PermissionDenied
    if log.status != LogImporStandar.Status.PRATINJAU or not log.file:
        raise GalatImpor("Impor ini sudah diproses.")
    hasil = impor_standar(path=log.file.path, nonaktifkan_yang_hilang=log.nonaktifkan_yang_hilang)
    _isi_log(log, hasil)
    log.status = (
        LogImporStandar.Status.BERHASIL if hasil.berhasil else LogImporStandar.Status.GAGAL_VALIDASI
    )
    log.selesai_pada = timezone.now()
    _hapus_file_sementara(log)
    log.save()
    return log


def batalkan_impor(*, log: LogImporStandar, user) -> LogImporStandar:
    if not bisa_proses_log_impor(user, log):
        raise PermissionDenied
    if log.status != LogImporStandar.Status.PRATINJAU:
        raise GalatImpor("Impor ini sudah diproses.")
    log.status = LogImporStandar.Status.DIBATALKAN
    log.selesai_pada = timezone.now()
    _hapus_file_sementara(log)
    log.save()
    return log


def impor_standar_cli(*, path, dry_run=False, nonaktifkan_yang_hilang=False) -> HasilImpor:
    hasil = impor_standar(
        path=path, dry_run=dry_run, nonaktifkan_yang_hilang=nonaktifkan_yang_hilang
    )
    if not dry_run:
        log = LogImporStandar(
            sumber=LogImporStandar.Sumber.CLI,
            nama_file=Path(path).name[:255],
            nonaktifkan_yang_hilang=nonaktifkan_yang_hilang,
            status=(
                LogImporStandar.Status.BERHASIL
                if hasil.berhasil
                else LogImporStandar.Status.GAGAL_VALIDASI
            ),
            selesai_pada=timezone.now(),
        )
        _isi_log(log, hasil)
        log.save()
    return hasil


# ---------------------------------------------------------------------------
# Template Excel
# ---------------------------------------------------------------------------


PETUNJUK = [
    (
        "Sheet Bab (opsional)",
        "Satu baris per bab. Kolom wajib: kode_kelompok, nama_kelompok, "
        "kode_bab, nama_bab. kode_pokja opsional; pokja baru dibuat otomatis bila belum ada.",
    ),
    (
        "Sheet EP (wajib)",
        "Satu baris per elemen penilaian. Kolom wajib: kode_bab, kode_standar, "
        "kode_ep, uraian_ep, jenis_bukti.",
    ),
    (
        "uraian_standar",
        "Cukup diisi sekali per standar; baris berikutnya dengan kode standar "
        "sama boleh dikosongkan.",
    ),
    (
        "jenis_bukti",
        "Dipisah koma. Hanya R (Regulasi), D (Dokumen), O (Observasi), "
        "W (Wawancara), S (Simulasi).",
    ),
    ("urutan", "Opsional, bilangan bulat. Bila kosong mengikuti urutan baris."),
    ("Kode", "Kode bab, standar, dan EP disimpan apa adanya sesuai dokumen resmi."),
]


def buat_template_workbook() -> Workbook:
    wb = Workbook()
    ws_ep = wb.active
    ws_ep.title = SHEET_EP
    ws_bab = wb.create_sheet(SHEET_BAB)
    ws_petunjuk = wb.create_sheet("Petunjuk")
    tebal = Font(bold=True)
    for ws, kolom in ((ws_ep, KOLOM_EP), (ws_bab, KOLOM_BAB)):
        ws.append(kolom)
        for sel in ws[1]:
            sel.font = tebal
        for i, nama in enumerate(kolom, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(
                14, len(nama) + 4
            )
    ws_petunjuk.append(["Bagian", "Keterangan"])
    for sel in ws_petunjuk[1]:
        sel.font = tebal
    for baris in PETUNJUK:
        ws_petunjuk.append(list(baris))
    ws_petunjuk.column_dimensions["A"].width = 22
    ws_petunjuk.column_dimensions["B"].width = 100
    return wb
