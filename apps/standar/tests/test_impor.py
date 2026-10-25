import pytest
from django.core.management import CommandError, call_command
from openpyxl import load_workbook

from apps.akun.models import Pokja
from apps.standar.models import Bab, ElemenPenilaian, ItemWawancara, LogImporStandar, Standar
from apps.standar.services import impor_standar

from .conftest import EP_CONTOH

pytestmark = pytest.mark.django_db


def test_impor_baru_membuat_bab_standar_ep(buat_excel):
    hasil = impor_standar(path=buat_excel())

    assert hasil.berhasil, hasil.galat
    bab = Bab.objects.get(kode="TKRS")
    assert bab.kelompok.kode == "KMRS"
    assert bab.pokja.kode == "POKJA-TKRS"
    assert Standar.objects.filter(bab=bab).count() == 2
    ep_b = ElemenPenilaian.objects.get(standar__kode="TKRS 1", kode="b")
    assert ep_b.jenis_bukti == ["D", "W"]
    # uraian_standar diambil dari baris sebelumnya dengan kode standar sama
    assert ep_b.standar.uraian == "Uraian standar contoh 1"
    assert ElemenPenilaian.objects.get(standar__kode="TKRS 2").jenis_bukti == ["R", "O"]
    assert hasil.ringkasan["ep"]["dibuat"] == 3


def test_impor_ulang_idempoten(buat_excel):
    path = buat_excel()
    impor_standar(path=path)
    hasil = impor_standar(path=path)

    assert hasil.berhasil
    assert ElemenPenilaian.objects.count() == 3
    assert Standar.objects.count() == 2
    assert hasil.ringkasan["ep"] == {"dibuat": 0, "diperbarui": 0, "tetap": 3}


def test_uraian_diubah_terperbarui_id_tetap(buat_excel):
    impor_standar(path=buat_excel())
    ep_lama = ElemenPenilaian.objects.get(standar__kode="TKRS 1", kode="a")

    ep_baru = [row.copy() for row in EP_CONTOH]
    ep_baru[0][5] = "Uraian EP contoh a (revisi)"
    hasil = impor_standar(path=buat_excel(ep=ep_baru))

    ep = ElemenPenilaian.objects.get(standar__kode="TKRS 1", kode="a")
    assert ep.pk == ep_lama.pk
    assert ep.uraian == "Uraian EP contoh a (revisi)"
    assert hasil.ringkasan["ep"]["diperbarui"] == 1


def test_jenis_bukti_tidak_valid_error_dengan_nomor_baris(buat_excel):
    ep = [row.copy() for row in EP_CONTOH]
    ep[1][6] = "D,X"
    hasil = impor_standar(path=buat_excel(ep=ep))

    assert not hasil.berhasil
    assert any("baris 3" in g and "'X'" in g for g in hasil.galat)
    assert not Bab.objects.exists()
    assert not ElemenPenilaian.objects.exists()


def test_kolom_wajib_kosong_dan_bab_tidak_dikenal(buat_excel):
    ep = [
        ["TKRS", "TKRS 1", "Uraian standar contoh", "", "a", "", "R", 1],
        ["XYZ", "XYZ 1", "Uraian standar contoh", "", "a", "Uraian EP contoh", "R", 1],
    ]
    hasil = impor_standar(path=buat_excel(ep=ep))

    assert any("baris 2" in g and "uraian_ep" in g for g in hasil.galat)
    assert any("baris 3" in g and "'XYZ'" in g for g in hasil.galat)
    assert not Bab.objects.exists()


def test_ep_duplikat_ditolak(buat_excel):
    ep = [EP_CONTOH[0], EP_CONTOH[0]]
    hasil = impor_standar(path=buat_excel(ep=ep))
    assert any("duplikat" in g for g in hasil.galat)


def test_header_tanpa_peduli_huruf_dan_spasi(buat_excel):
    header = [
        "Kode Bab",
        "KODE_STANDAR",
        "uraian standar",
        "Maksud Tujuan",
        "kode ep",
        "Uraian EP",
        "Jenis Bukti",
        "Urutan",
    ]
    hasil = impor_standar(path=buat_excel(header_ep=header))
    assert hasil.berhasil, hasil.galat


def test_ep_hilang_dinonaktifkan_bukti_tetap_ada(buat_excel):
    impor_standar(path=buat_excel())
    ep_b = ElemenPenilaian.objects.get(standar__kode="TKRS 1", kode="b")
    ItemWawancara.objects.create(elemen_penilaian=ep_b, narasumber="Jabatan contoh")

    hasil = impor_standar(
        path=buat_excel(ep=[EP_CONTOH[0], EP_CONTOH[2]]), nonaktifkan_yang_hilang=True
    )

    assert hasil.berhasil
    assert hasil.dinonaktifkan == ["TKRS 1 EP b"]
    assert not ElemenPenilaian.objects.filter(pk=ep_b.pk).exists()
    assert ElemenPenilaian.semua.get(pk=ep_b.pk).dihapus_pada is not None
    assert ItemWawancara.objects.filter(elemen_penilaian_id=ep_b.pk).exists()


def test_ep_hilang_tanpa_flag_tidak_dinonaktifkan(buat_excel):
    impor_standar(path=buat_excel())
    impor_standar(path=buat_excel(ep=[EP_CONTOH[0]]))
    assert ElemenPenilaian.objects.count() == 3


def test_ep_nonaktif_dipulihkan_saat_muncul_lagi(buat_excel):
    impor_standar(path=buat_excel())
    impor_standar(path=buat_excel(ep=[EP_CONTOH[0]]), nonaktifkan_yang_hilang=True)
    assert ElemenPenilaian.objects.count() == 1
    hasil = impor_standar(path=buat_excel())
    assert ElemenPenilaian.objects.count() == 3
    assert ElemenPenilaian.semua.count() == 3
    assert hasil.ringkasan["ep"]["diperbarui"] == 2


def test_dry_run_tidak_menyimpan(buat_excel):
    hasil = impor_standar(path=buat_excel(), dry_run=True)
    assert hasil.berhasil
    assert hasil.ringkasan["ep"]["dibuat"] == 3
    assert not Bab.objects.exists()
    assert not Pokja.objects.exists()


def test_tanpa_sheet_bab_memakai_bab_yang_ada(buat_excel):
    impor_standar(path=buat_excel())
    ep = [["TKRS", "TKRS 3", "Uraian standar contoh 3", "", "a", "Uraian EP contoh", "S", 1]]
    hasil = impor_standar(path=buat_excel(ep=ep, bab=None))
    assert hasil.berhasil, hasil.galat
    assert Standar.objects.filter(bab__kode="TKRS").count() == 3


def test_file_bukan_excel(tmp_path):
    path = tmp_path / "rusak.xlsx"
    path.write_bytes(b"bukan excel")
    hasil = impor_standar(path=path)
    assert hasil.galat == ["File tidak dapat dibuka sebagai Excel (.xlsx)."]


def test_command_dry_run_dan_simpan(buat_excel, capsys):
    path = buat_excel()
    call_command("impor_standar", str(path), "--dry-run")
    assert not Bab.objects.exists()
    assert not LogImporStandar.objects.exists()

    call_command("impor_standar", str(path))
    assert ElemenPenilaian.objects.count() == 3
    log = LogImporStandar.objects.get()
    assert log.sumber == LogImporStandar.Sumber.CLI
    assert log.status == LogImporStandar.Status.BERHASIL


def test_command_error_keluar_dengan_galat(buat_excel):
    ep = [row.copy() for row in EP_CONTOH]
    ep[0][6] = "Z"
    with pytest.raises(CommandError):
        call_command("impor_standar", str(buat_excel(ep=ep)))
    assert not Bab.objects.exists()


def test_command_buat_template(tmp_path):
    path = tmp_path / "template.xlsx"
    call_command("buat_template_standar", str(path))
    wb = load_workbook(path)
    assert {"EP", "Bab", "Petunjuk"} <= set(wb.sheetnames)
    assert [c.value for c in wb["EP"][1]][:2] == ["kode_bab", "kode_standar"]
